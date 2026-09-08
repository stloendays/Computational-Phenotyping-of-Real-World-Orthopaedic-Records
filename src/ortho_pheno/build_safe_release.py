#!/usr/bin/env python3
"""Build privacy-preserving aggregate research artifacts from local orthopaedic exports.

IMPORTANT: this script reads raw local clinical files but writes aggregate outputs only.
It intentionally never writes patient-level rows, direct identifiers, free-text notes, or dates.
"""
from pathlib import Path
from collections import defaultdict, Counter
from datetime import datetime, timedelta
import csv, math, re, statistics, sys
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
from legacy_xls import parse_biff
from rules import is_hallux_valgus, is_first_cmc_oa, is_wrist_scaphoid, SCAPHOID_CHRONIC_RE

DIRECT_IDS = {'姓名','住院id','住院流水号','住院号','病案号','电话','联系人电话','就诊卡号','申请单ID'}
DISEASES = {
    'hallux_valgus': '拇外翻',
    'first_cmc_oa': '腕掌关节炎',
    'scaphoid_fracture': '舟骨骨折',
}

PROC_PATTERNS = {
    'hallux_valgus': {
        'osteotomy': r'截骨', 'chevron': r'Chevron|chevron', 'akin': r'Akin|AKIN|akin',
        'scarf': r'Scarf|SCARF|scarf', 'fusion': r'融合', 'k_wire': r'克氏针|钢针',
        'soft_tissue': r'肌腱|韧带|关节囊|松解', 'resection': r'切除',
    },
    'first_cmc_oa': {
        'trapeziectomy': r'大多角骨.*切除|切除.*大多角骨', 'tendon_procedure': r'肌腱',
        'ligament_procedure': r'韧带', 'arthroplasty': r'关节成形|关节置换', 'fusion': r'融合',
    },
    'scaphoid_fracture': {
        'internal_fixation': r'内固定|螺钉|空心钉|钢针', 'bone_graft': r'植骨|取骨',
        'fusion': r'融合', 'debridement': r'清创|病灶清除', 'hardware_removal': r'取出内固定|内固定.*去除',
    },
}

def norm_id(v):
    if v is None: return None
    s=str(v).strip()
    if s.endswith('.0') and s[:-2].isdigit(): s=s[:-2]
    return s or None

def excel_dt(v):
    if isinstance(v, datetime): return v
    if isinstance(v, (int,float)) and not isinstance(v,bool):
        try: return datetime(1899,12,30)+timedelta(days=float(v))
        except Exception: return None
    if isinstance(v,str):
        s=v.strip()
        for fmt in ('%Y-%m-%d %H:%M:%S','%Y-%m-%d','%Y/%m/%d %H:%M:%S','%Y/%m/%d'):
            try: return datetime.strptime(s,fmt)
            except Exception: pass
    return None

def read_xlsx(path):
    wb=load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        headers=[str(ws.cell(1,c).value).strip() if ws.cell(1,c).value is not None else '' for c in range(1,ws.max_column+1)]
        rows=[]
        for vals in ws.iter_rows(min_row=2,values_only=True):
            rows.append({headers[i]:vals[i] for i in range(len(headers)) if headers[i]})
        yield ws.title, headers, rows

def classify_file(name):
    if name.endswith('检查.xls'): return 'imaging_exam'
    if name.endswith('检验.xls'): return 'laboratory'
    if '主诉和专科查体' in name: return 'complaint_exam'
    if '基本信息和诊断信息' in name: return 'demographics_diagnoses'
    if '能查到手术内容' in name: return 'detailed_operating_note'
    if '无法查询手术内容' in name or '能无法查询手术内容' in name: return 'surgery_name_only'
    return 'other'

def disease_key_from_name(name):
    for k,zh in DISEASES.items():
        if zh in name or (k=='hallux_valgus' and name.startswith('拇外翻')) or (k=='first_cmc_oa' and name.startswith('腕掌关节炎')) or (k=='scaphoid_fracture' and name.startswith('舟骨骨折')):
            return k
    return 'unknown'

def qtile(vals,p):
    vals=sorted(float(x) for x in vals if x is not None and not math.isnan(float(x)))
    if not vals:return None
    pos=(len(vals)-1)*p; lo=math.floor(pos); hi=math.ceil(pos)
    return vals[lo] if lo==hi else vals[lo]+(vals[hi]-vals[lo])*(pos-lo)

def fmt(v, nd=1): return '' if v is None else f'{v:.{nd}f}'
def suppress(n): return '<5' if isinstance(n,int) and 0<n<5 else str(n)

def build(raw_dir, out_dir):
    raw_dir=Path(raw_dir); out_dir=Path(out_dir); (out_dir/'data/aggregate').mkdir(parents=True,exist_ok=True); (out_dir/'data/schema').mkdir(parents=True,exist_ok=True)
    inventory=[]; recs={k:defaultdict(lambda:{'texts':[],'clinical_texts':[],'sex':None,'age':None,'height':None,'weight':None,'admit':None,'discharge':None,'complaint':False,'exam':False,'surgery_any':False,'surgery_detail':False,'surgery_texts':[]}) for k in DISEASES}
    # xlsx
    for p in sorted(raw_dir.glob('*.xlsx')):
        dk=disease_key_from_name(p.name); module=classify_file(p.name); sheets=[]; total_rows=0; fields=set(); ids=set()
        for s,h,rows in read_xlsx(p):
            sheets.append(s); total_rows += len(rows); fields.update(h)
            for d in rows:
                pid=norm_id(d.get('住院id'))
                if pid: ids.add(pid)
                if dk=='unknown' or not pid: continue
                r=recs[dk][pid]
                for key,val in d.items():
                    if val is None or key in DIRECT_IDS: continue
                    if isinstance(val,str) and val.strip(): r['texts'].append(val.strip())
                for key in ('主要诊断名称','主要诊断描述','其他诊断名称','其他诊断描述','主诉','专科查体'):
                    val=d.get(key)
                    if isinstance(val,str) and val.strip(): r['clinical_texts'].append(val.strip())
                r['sex']=r['sex'] or d.get('性别'); r['age']=r['age'] if r['age'] is not None else d.get('年龄'); r['height']=r['height'] if r['height'] is not None else d.get('身高'); r['weight']=r['weight'] if r['weight'] is not None else d.get('体重')
                for key in ('入院日期','入院时间'):
                    dt=excel_dt(d.get(key));
                    if dt and (r['admit'] is None or dt<r['admit']): r['admit']=dt
                for key in ('出院日期','出院时间'):
                    dt=excel_dt(d.get(key));
                    if dt and (r['discharge'] is None or dt>r['discharge']): r['discharge']=dt
                if d.get('主诉'): r['complaint']=True
                if d.get('专科查体'): r['exam']=True
                if d.get('手术时间') or d.get('手术名称') or d.get('手术记录内容'): r['surgery_any']=True
                if d.get('手术记录内容'):
                    r['surgery_detail']=True; r['surgery_texts'].append(str(d.get('手术记录内容')))
                if d.get('手术名称'): r['surgery_texts'].append(str(d.get('手术名称')))
        inventory.append({'disease':dk,'module':module,'source_file':p.name,'format':'xlsx','sheets':'|'.join(sheets),'data_rows':total_rows,'columns':len(fields),'unique_admission_keys':len(ids),'contains_direct_identifiers':'yes','public_release':'no'})
    # strict sets
    strict={}
    for dk,rs in recs.items():
        s=set()
        for pid,r in rs.items():
            t=' '.join(r['texts'])
            if dk=='hallux_valgus' and is_hallux_valgus(t): s.add(pid)
            elif dk=='first_cmc_oa' and is_first_cmc_oa(t): s.add(pid)
            elif dk=='scaphoid_fracture' and is_wrist_scaphoid(t): s.add(pid)
        strict[dk]=s
    # legacy xls inventory + coverage sets
    lab_cov={k:defaultdict(set) for k in DISEASES}; exam_cov={k:set() for k in DISEASES}
    for p in sorted(raw_dir.glob('*.xls')):
        dk=disease_key_from_name(p.name); module=classify_file(p.name); data=parse_biff(p); sheets=[]; total_rows=0; colmax=0; ids=set()
        for s,rows in data.items():
            sheets.append(s); total_rows += max(0,len(rows)-1); colmax=max(colmax,len(rows[0]) if rows else 0)
            if not rows: continue
            h=rows[0]; idx={x:i for i,x in enumerate(h)}
            idc=idx.get('住院流水号')
            for row in rows[1:]:
                pid=norm_id(row[idc]) if idc is not None and idc<len(row) else None
                if not pid: continue
                ids.add(pid)
                if module=='imaging_exam': exam_cov[dk].add(pid)
                elif module=='laboratory':
                    if '凝血' in s: cat='coagulation'
                    elif '血常规' in s: cat='cbc'
                    elif '红细胞沉' in s: cat='esr'
                    elif 'C反应蛋白' in s: cat='crp'
                    elif '肝功' in s: cat='liver_function'
                    elif '肾功' in s: cat='kidney_function'
                    elif '离子' in s: cat='electrolytes'
                    else: cat='other'
                    lab_cov[dk][cat].add(pid)
        inventory.append({'disease':dk,'module':module,'source_file':p.name,'format':'xls','sheets':'|'.join(sheets),'data_rows':total_rows,'columns':colmax,'unique_admission_keys':len(ids),'contains_direct_identifiers':'yes','public_release':'no'})
    # inventory CSV
    fields=['disease','module','source_file','format','sheets','data_rows','columns','unique_admission_keys','contains_direct_identifiers','public_release']
    with open(out_dir/'data/schema/source_inventory.csv','w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fieldnames=fields); w.writeheader(); w.writerows(inventory)
    # field map
    fmap=[
        ('admission_key','住院id / 住院流水号','linkage only','never public'),('sex','性别','structured baseline','aggregate only'),('age','年龄','structured baseline','aggregate only'),('height','身高','structured baseline','aggregate only'),('weight','体重','structured baseline','aggregate only'),('admission_date','入院日期/入院时间','temporal alignment','year-level aggregate only'),('discharge_date','出院日期/出院时间','temporal alignment','not public row-level'),('main_diagnosis','主要诊断名称/描述','phenotype source','derived aggregate only'),('other_diagnosis','其他诊断名称/描述','phenotype source','derived aggregate only'),('complaint','主诉','NLP phenotype source','never public raw text'),('physical_exam','专科查体','NLP phenotype source','never public raw text'),('operation','手术名称/手术记录内容','procedure phenotype source','never public raw text'),('exam_report','检查项目/所见/结果','exam phenotype source','never public raw text'),('laboratory','项目/检验项目/结果','laboratory coverage','aggregate only')]
    with open(out_dir/'data/schema/field_map.csv','w',encoding='utf-8-sig',newline='') as f:
        w=csv.writer(f); w.writerow(['canonical_field','source_field','research_role','public_release']); w.writerows(fmap)
    # cohort overview
    rows=[]
    for dk,zh in DISEASES.items():
        rs=recs[dk]; ss=strict[dk]; ages=[]; female=0; sexn=0; bmis=[]
        for pid in ss:
            r=rs[pid]
            try:
                a=float(r['age']);
                if 0<a<120: ages.append(a)
            except Exception: pass
            if r['sex'] is not None:
                sexn+=1; female += 1 if str(r['sex']).strip() in ('女','女性','F','Female') else 0
            try:
                h=float(r['height']); w=float(r['weight']);
                if h>3: h/=100
                bmi=w/(h*h)
                if 10<bmi<60:bmis.append(bmi)
            except Exception: pass
        rows.append({
            'disease':dk,'candidate_admissions':len(rs),'strict_phenotype_admissions':len(ss),
            'female_percent':fmt(100*female/sexn if sexn else None),'age_median':fmt(statistics.median(ages) if ages else None),'age_q1':fmt(qtile(ages,.25)),'age_q3':fmt(qtile(ages,.75)),'bmi_median':fmt(statistics.median(bmis) if bmis else None),
            'with_complaint':sum(rs[p]['complaint'] for p in ss),'with_physical_exam':sum(rs[p]['exam'] for p in ss),'with_any_surgery_record':sum(rs[p]['surgery_any'] for p in ss),'with_detailed_operating_note':sum(rs[p]['surgery_detail'] for p in ss),
            'with_any_exam_record':len(ss & exam_cov[dk]),'with_cbc':len(ss & lab_cov[dk]['cbc']),'with_coagulation':len(ss & lab_cov[dk]['coagulation']),'with_esr':len(ss & lab_cov[dk]['esr']),'with_crp':len(ss & lab_cov[dk]['crp']),'with_liver_function':len(ss & lab_cov[dk]['liver_function']),'with_kidney_function':len(ss & lab_cov[dk]['kidney_function']),'with_electrolytes':len(ss & lab_cov[dk]['electrolytes'])
        })
    of=list(rows[0].keys())
    with open(out_dir/'data/aggregate/cohort_overview.csv','w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fieldnames=of); w.writeheader(); w.writerows(rows)
    # years, suppressed
    yrs=[]
    for dk in DISEASES:
        c=Counter(recs[dk][p]['admit'].year for p in strict[dk] if recs[dk][p]['admit'])
        for y in range(2015,2026): yrs.append({'disease':dk,'year':y,'strict_admissions':suppress(c.get(y,0))})
    with open(out_dir/'data/aggregate/year_distribution.csv','w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fieldnames=['disease','year','strict_admissions']); w.writeheader(); w.writerows(yrs)
    # procedure counts among strict cases with detailed note
    proc=[]
    for dk,pats in PROC_PATTERNS.items():
        denom=sum(recs[dk][p]['surgery_detail'] for p in strict[dk])
        for label,pat in pats.items():
            n=0
            rg=re.compile(pat,re.I)
            for pid in strict[dk]:
                if recs[dk][pid]['surgery_detail'] and rg.search(' '.join(recs[dk][pid]['surgery_texts'])): n+=1
            proc.append({'disease':dk,'procedure_phenotype':label,'n':suppress(n),'denominator_detailed_notes':denom,'percent':('<5' if 0<n<5 else fmt(100*n/denom if denom else None))})
    with open(out_dir/'data/aggregate/procedure_phenotype_counts.csv','w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fieldnames=['disease','procedure_phenotype','n','denominator_detailed_notes','percent']); w.writeheader(); w.writerows(proc)
    # scaphoid chronic phenotype aggregate: non-operative clinical text only
    chronic=sum(1 for pid in strict['scaphoid_fracture'] if SCAPHOID_CHRONIC_RE.search(' '.join(recs['scaphoid_fracture'][pid]['clinical_texts'])))
    other=len(strict['scaphoid_fracture'])-chronic
    with open(out_dir/'data/aggregate/scaphoid_phenotype_groups.csv','w',encoding='utf-8-sig',newline='') as f:
        w=csv.writer(f); w.writerow(['phenotype_group','n','interpretation']); w.writerow(['established_chronic_or_nonunion',chronic,'Text-supported established chronic/nonunion phenotype defined from non-operative clinical text; not necessarily incident nonunion observed longitudinally.']); w.writerow(['other_wrist_scaphoid',other,'Wrist-scaphoid cases without the v0.2 chronic/nonunion terms in diagnosis/complaint/examination text.'])

if __name__=='__main__':
    import argparse
    ap=argparse.ArgumentParser(); ap.add_argument('--raw-dir',required=True); ap.add_argument('--out-dir',required=True); args=ap.parse_args(); build(args.raw_dir,args.out_dir)
