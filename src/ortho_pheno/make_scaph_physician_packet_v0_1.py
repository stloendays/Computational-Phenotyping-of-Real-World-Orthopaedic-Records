#!/usr/bin/env python3
"""Generate LOCAL-ONLY two-stage physician review packets for the scaphoid paper.

Why two stages?
---------------
The clinical cohort must not be locked by the deterministic anatomy rule before
physician review. Stage 1 therefore reviews anatomy in all 88 broad scaphoid
candidates. Stage 2 is generated only after Stage-1 anatomy labels are complete
and selects *physician-confirmed* wrist-scaphoid records for state/procedure
review. This avoids verification bias from reviewing only deterministic positives.

Stage 2 also captures a physician-adjudicated wrist-related duration variable for
a prespecified exploratory secondary hypothesis. Automated duration-parser output
is never shown to the reviewer.

Patient-level outputs contain protected clinical text and pseudonymous study IDs.
They MUST remain local and are intended for ``annotations/private/`` only.

The pseudonym salt is supplied locally and is never committed to the repository.
"""
from __future__ import annotations

from collections import defaultdict
from pathlib import Path
import argparse
import csv
import hashlib
import os
import re

from openpyxl import load_workbook

SCAPHOID_FILES = (
    '2015-2025舟骨骨折患者基本信息和诊断信息.xlsx',
    '2015-2025舟骨骨折患者主诉和专科查体.xlsx',
    '2015-2025舟骨骨折患者无法查询手术内容.xlsx',
    '2015-2025舟骨骨折患者能查到手术内容的患者.xlsx',
)

ANATOMY_ALLOWED = {'wrist_scaphoid','foot_navicular','other','uncertain'}

PHONE_RE = re.compile(r'(?<!\d)(?:\+?86[- ]?)?1[3-9]\d{9}(?!\d)')
CN_ID_RE = re.compile(r'(?<![0-9A-Za-z])\d{17}[0-9Xx](?![0-9A-Za-z])')
LONG_NUMBER_RE = re.compile(r'(?<!\d)\d{9,}(?!\d)')


def norm_id(value):
    if value is None:
        return None
    s = str(value).strip()
    if s.endswith('.0') and s[:-2].isdigit():
        s = s[:-2]
    return s or None


def load_salt(path: str | None) -> str:
    if path:
        salt = Path(path).read_text(encoding='utf-8').strip()
    else:
        salt = os.environ.get('ORTHO_ANNOTATION_SALT','').strip()
    if len(salt) < 16:
        raise ValueError('Provide a local pseudonym salt of at least 16 characters via --salt-file or ORTHO_ANNOTATION_SALT.')
    return salt


def make_study_id(admission_key: str, salt: str) -> str:
    digest = hashlib.sha256(f'{salt}|scaphoid|{admission_key}'.encode('utf-8')).hexdigest()
    return f'SCA-{digest[:14]}'


def review_score(layer: str, study_id: str, salt: str) -> float:
    digest = hashlib.sha256(f'{salt}|double-review|{layer}|{study_id}'.encode('utf-8')).hexdigest()
    return int(digest[:16],16) / float(16**16 - 1)


def clean_text(value, patient_name=None):
    if value is None:
        return ''
    text = str(value).strip()
    if not text:
        return ''
    if patient_name:
        name = str(patient_name).strip()
        if name:
            text = text.replace(name,'[NAME]')
    text = PHONE_RE.sub('[PHONE]',text)
    text = CN_ID_RE.sub('[ID]',text)
    text = LONG_NUMBER_RE.sub('[LONG_NUMBER]',text)
    text = re.sub(r'\s+',' ',text).strip()
    return text


def read_xlsx(path: Path):
    wb = load_workbook(path,read_only=True,data_only=True)
    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)
        try:
            header = [str(x).strip() if x is not None else '' for x in next(rows)]
        except StopIteration:
            continue
        for values in rows:
            yield {header[i]:values[i] for i in range(min(len(header),len(values))) if header[i]}


def empty_record():
    return {
        'name':None,
        'diagnosis':[],
        'complaint':[],
        'exam':[],
        'operation_name':[],
        'operation_note':[],
    }


def add_unique(items,value):
    if value and value not in items:
        items.append(value)


def collect(raw_dir: Path):
    records = defaultdict(empty_record)
    for filename in SCAPHOID_FILES:
        path = raw_dir / filename
        if not path.exists():
            raise FileNotFoundError(path)
        for row in read_xlsx(path):
            key = norm_id(row.get('住院id'))
            if not key:
                continue
            rec = records[key]
            if not rec['name'] and row.get('姓名'):
                rec['name'] = str(row.get('姓名')).strip()
            for field in ('主要诊断名称','主要诊断描述','其他诊断名称','其他诊断描述'):
                add_unique(rec['diagnosis'],clean_text(row.get(field),rec['name']))
            add_unique(rec['complaint'],clean_text(row.get('主诉'),rec['name']))
            add_unique(rec['exam'],clean_text(row.get('专科查体'),rec['name']))
            add_unique(rec['operation_name'],clean_text(row.get('手术名称'),rec['name']))
            add_unique(rec['operation_note'],clean_text(row.get('手术记录内容'),rec['name']))
    return records


def join(items):
    return ' || '.join(x for x in items if x)


def write_csv(path: Path, rows, fields):
    path.parent.mkdir(parents=True,exist_ok=True)
    with path.open('w',encoding='utf-8-sig',newline='') as f:
        writer = csv.DictWriter(f,fieldnames=fields)
        writer.writeheader(); writer.writerows(rows)


def choose_double_review(rows, layer: str, salt: str, fraction: float):
    scored = sorted((review_score(layer,row['study_id'],salt),row['study_id']) for row in rows)
    if not scored:
        return []
    n = max(1,round(fraction*len(scored)))
    return [{'study_id':sid,'layer':layer,'hash_score':f'{score:.8f}'} for score,sid in scored[:n]]


def build_stage1(raw_dir: Path, output_dir: Path, salt: str, fraction: float):
    """Build the frozen Stage-1 anatomy packet.

    Keep this output schema stable: the pre-annotation Stage-1 packet has already
    been frozen by SHA-256. Stage-2 extensions must not change Stage-1 contents.
    """
    records = collect(raw_dir)
    rows = []
    for key,rec in records.items():
        rows.append({
            'study_id':make_study_id(key,salt),
            'diagnosis_text':join(rec['diagnosis']),
            'complaint_text':join(rec['complaint']),
            'physical_exam_text':join(rec['exam']),
            'operation_name_text':join(rec['operation_name']),
            'operation_note_text':join(rec['operation_note']),
            'gold_anatomy_label':'',
            'reviewer_confidence':'',
            'reviewer_comment':'',
        })
    rows.sort(key=lambda x:x['study_id'])
    if len(rows) != 88:
        raise AssertionError(f'expected 88 broad scaphoid candidates, found {len(rows)}')

    fields = ['study_id','diagnosis_text','complaint_text','physical_exam_text','operation_name_text','operation_note_text','gold_anatomy_label','reviewer_confidence','reviewer_comment']
    write_csv(output_dir/'scaphoid_anatomy_review.csv',rows,fields)
    doubles = choose_double_review(rows,'anatomy',salt,fraction)
    write_csv(output_dir/'scaphoid_anatomy_double_review_manifest.csv',doubles,['study_id','layer','hash_score'])
    return {'anatomy_primary':len(rows),'anatomy_double_review':len(doubles)}


def read_anatomy_gold(path: Path):
    with path.open(encoding='utf-8-sig',newline='') as f:
        rows = list(csv.DictReader(f))
    if len(rows) != 88:
        raise ValueError(f'anatomy gold must contain 88 rows; found {len(rows)}')
    result = {}
    for i,row in enumerate(rows,2):
        sid = row.get('study_id','').strip()
        label = row.get('gold_anatomy_label','').strip()
        if not sid:
            raise ValueError(f'blank study_id at row {i}')
        if sid in result:
            raise ValueError(f'duplicate study_id: {sid}')
        if label not in ANATOMY_ALLOWED:
            raise ValueError(f'row {i}: invalid or blank anatomy label {label!r}')
        result[sid] = label
    return result


def build_stage2(raw_dir: Path, output_dir: Path, salt: str, anatomy_gold_path: Path, fraction: float):
    records = collect(raw_dir)
    gold = read_anatomy_gold(anatomy_gold_path)

    by_sid = {make_study_id(key,salt):rec for key,rec in records.items()}
    if set(gold) != set(by_sid):
        missing = sorted(set(by_sid)-set(gold))
        extra = sorted(set(gold)-set(by_sid))
        raise ValueError(f'anatomy gold IDs do not match regenerated candidate set; missing={len(missing)}, extra={len(extra)}')

    wrist_ids = sorted(sid for sid,label in gold.items() if label == 'wrist_scaphoid')
    state_rows = []
    procedure_rows = []
    for sid in wrist_ids:
        rec = by_sid[sid]
        state_rows.append({
            'study_id':sid,
            'diagnosis_text':join(rec['diagnosis']),
            'complaint_text':join(rec['complaint']),
            'physical_exam_text':join(rec['exam']),
            'gold_scaphoid_state':'',
            'gold_relevant_duration_present':'',
            'gold_relevant_duration_value':'',
            'gold_relevant_duration_unit':'',
            'gold_duration_basis':'',
            'reviewer_confidence':'',
            'reviewer_comment':'',
        })
        if rec['operation_note']:
            procedure_rows.append({
                'study_id':sid,
                'operation_name_text':join(rec['operation_name']),
                'operation_note_text':join(rec['operation_note']),
                'gold_target_disease_procedure_present':'',
                'gold_internal_fixation':'',
                'gold_bone_graft':'',
                'gold_reconstruction':'',
                'gold_fusion':'',
                'reviewer_confidence':'',
                'reviewer_comment':'',
            })

    write_csv(output_dir/'scaphoid_state_review.csv',state_rows,
              ['study_id','diagnosis_text','complaint_text','physical_exam_text','gold_scaphoid_state','gold_relevant_duration_present','gold_relevant_duration_value','gold_relevant_duration_unit','gold_duration_basis','reviewer_confidence','reviewer_comment'])
    write_csv(output_dir/'scaphoid_procedure_review.csv',procedure_rows,
              ['study_id','operation_name_text','operation_note_text','gold_target_disease_procedure_present','gold_internal_fixation','gold_bone_graft','gold_reconstruction','gold_fusion','reviewer_confidence','reviewer_comment'])

    doubles = choose_double_review(state_rows,'state',salt,fraction) + choose_double_review(procedure_rows,'procedure',salt,fraction)
    write_csv(output_dir/'scaphoid_downstream_double_review_manifest.csv',doubles,['study_id','layer','hash_score'])
    return {
        'physician_confirmed_wrist_scaphoid':len(state_rows),
        'wrist_with_detailed_note_module':len(procedure_rows),
        'downstream_double_review':len(doubles),
        'duration_fields_added_to_state_review':4,
    }


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--raw-dir',required=True)
    ap.add_argument('--output-dir',required=True)
    ap.add_argument('--salt-file')
    ap.add_argument('--double-review-fraction',type=float,default=0.20)
    sub = ap.add_subparsers(dest='stage',required=True)
    sub.add_parser('anatomy')
    p2 = sub.add_parser('downstream')
    p2.add_argument('--anatomy-gold',required=True)
    args = ap.parse_args()

    if not 0 < args.double_review_fraction <= 1:
        raise ValueError('--double-review-fraction must be in (0,1]')
    salt = load_salt(args.salt_file)
    raw_dir = Path(args.raw_dir)
    output_dir = Path(args.output_dir)
    if args.stage == 'anatomy':
        summary = build_stage1(raw_dir,output_dir,salt,args.double_review_fraction)
    else:
        summary = build_stage2(raw_dir,output_dir,salt,Path(args.anatomy_gold),args.double_review_fraction)
    for key,value in summary.items():
        print(f'{key}: {value}')


if __name__ == '__main__':
    main()
