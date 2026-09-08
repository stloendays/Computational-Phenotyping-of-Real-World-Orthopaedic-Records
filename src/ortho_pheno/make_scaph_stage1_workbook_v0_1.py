#!/usr/bin/env python3
"""Create a LOCAL-ONLY Excel workbook for Stage-1 scaphoid anatomy review.

Inputs are the private 88-row anatomy-review CSV and the prespecified Reviewer-2
manifest. The workbook contains clinical text and pseudonymous study IDs and must
remain private. No rule/model predictions are included.
"""
from __future__ import annotations

from pathlib import Path
import argparse
import csv

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

LABELS = ['wrist_scaphoid','foot_navicular','other','uncertain']
CONFIDENCE = ['high','medium','low']
FIELDS = [
    'study_id','diagnosis_text','complaint_text','physical_exam_text',
    'operation_name_text','operation_note_text','gold_anatomy_label',
    'reviewer_confidence','reviewer_comment',
]
HEADERS_ZH = {
    'study_id':'Study ID（伪匿名）',
    'diagnosis_text':'诊断文本',
    'complaint_text':'主诉',
    'physical_exam_text':'专科查体',
    'operation_name_text':'手术名称',
    'operation_note_text':'详细手术记录',
    'gold_anatomy_label':'解剖标签（医生填写）',
    'reviewer_confidence':'置信度（医生填写）',
    'reviewer_comment':'备注/判断依据（医生填写）',
}


def read_csv(path: Path):
    with path.open(encoding='utf-8-sig', newline='') as f:
        return list(csv.DictReader(f))


def style_review_sheet(ws, rows, reviewer2=False):
    dark='1F1F1F'; white='FFFFFF'; green='008000'; blue='0000FF'; gray='666666'
    green_light='E2F0D9'; blue_light='D9EAF7'; very_light='F7F7F7'
    med_dark=Side(style='medium',color=dark)
    ws.sheet_view.showGridLines=False
    ws.freeze_panes='A2'
    ws.auto_filter.ref=f'A1:I{len(rows)+1}'
    ws.row_dimensions[1].height=42
    for col,field in enumerate(FIELDS,1):
        c=ws.cell(1,col,HEADERS_ZH[field])
        c.font=Font(name='Arial',size=10,bold=True,color=white)
        c.fill=PatternFill('solid',fgColor=dark)
        c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
        c.border=Border(bottom=med_dark)
    widths=[20,44,42,42,34,70,25,18,42]
    for col,width in enumerate(widths,1):
        ws.column_dimensions[get_column_letter(col)].width=width
    label_dv=DataValidation(type='list',formula1="='_Lists'!$A$1:$A$4",allow_blank=True)
    conf_dv=DataValidation(type='list',formula1="='_Lists'!$B$1:$B$3",allow_blank=True)
    ws.add_data_validation(label_dv);ws.add_data_validation(conf_dv)
    label_dv.add(f'G2:G{len(rows)+1}');conf_dv.add(f'H2:H{len(rows)+1}')
    for r_idx,row in enumerate(rows,2):
        for c_idx,field in enumerate(FIELDS,1):
            c=ws.cell(r_idx,c_idx,row.get(field,''))
            c.alignment=Alignment(wrap_text=True,vertical='top')
            if c_idx in (2,3,4,5,6):
                c.font=Font(name='Arial',size=9,color=green)
                c.fill=PatternFill('solid',fgColor=green_light)
            elif c_idx in (7,8,9):
                c.font=Font(name='Arial',size=9,color=blue)
                c.fill=PatternFill('solid',fgColor=blue_light)
            else:
                c.font=Font(name='Arial',size=9,color=gray)
                c.fill=PatternFill('solid',fgColor=very_light)
        ws.row_dimensions[r_idx].height=92
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0;ws.page_orientation='landscape'
    ws.print_title_rows='1:1'
    ws.oddFooter.center.text='Confidential research annotation — do not upload publicly'
    ws['G1'].comment=Comment(
        'Reviewer 2 必须独立填写；不要查看 Reviewer 1 标签。' if reviewer2 else
        'Reviewer 1 独立填写；不要参考规则预测或 Reviewer 2。',
        'Research protocol'
    )


def build(primary_csv: Path, manifest_csv: Path, output: Path):
    primary=read_csv(primary_csv);manifest=read_csv(manifest_csv)
    if len(primary)!=88:
        raise ValueError(f'expected 88 primary rows, got {len(primary)}')
    if len(manifest)!=18:
        raise ValueError(f'expected 18 Stage-1 Reviewer-2 rows, got {len(manifest)}')
    by_id={r['study_id']:r for r in primary}
    if len(by_id)!=88:
        raise ValueError('duplicate primary study_id')
    r2=[]
    for m in manifest:
        sid=m.get('study_id','').strip()
        if sid not in by_id:
            raise ValueError(f'Reviewer-2 ID missing from primary: {sid}')
        row=dict(by_id[sid])
        row['gold_anatomy_label']='';row['reviewer_confidence']='';row['reviewer_comment']=''
        r2.append(row)

    wb=Workbook();info=wb.active;info.title='填写说明'
    r1=wb.create_sheet('Reviewer1_88');r2ws=wb.create_sheet('Reviewer2_18');lists=wb.create_sheet('_Lists')
    lists.sheet_state='hidden'
    for i,v in enumerate(LABELS,1):lists.cell(i,1,v)
    for i,v in enumerate(CONFIDENCE,1):lists.cell(i,2,v)

    info.sheet_view.showGridLines=False
    info.column_dimensions['B'].width=25;info.column_dimensions['C'].width=90
    info.merge_cells('B2:C2');info['B2']='腕舟骨研究 Stage 1：解剖部位人工复核'
    info['B2'].font=Font(name='Arial',size=18,bold=True,color='FFFFFF');info['B2'].fill=PatternFill('solid',fgColor='1F1F1F')
    rows=[
        ('用途','本工作簿只用于 Stage 1 解剖部位判定。请依据去标识化病历上下文独立判断。'),
        ('标签','wrist_scaphoid / foot_navicular / other / uncertain。证据不足时保留 uncertain。'),
        ('盲法','不要参考规则/模型预测、统计结果或另一位复核者答案。'),
        ('Reviewer 1','Reviewer1_88：填写全部88例。'),
        ('Reviewer 2','Reviewer2_18：仅填写预先冻结的18例子集，并与 Reviewer 1 独立。'),
        ('隐私','包含伪匿名 study_id 和临床文本，仍属敏感研究资料；不得上传公开 GitHub。'),
    ]
    for i,(k,v) in enumerate(rows,4):
        info.cell(i,2,k).font=Font(bold=True,color='666666')
        info.cell(i,3,v).alignment=Alignment(wrap_text=True,vertical='top')
        info.row_dimensions[i].height=38
    style_review_sheet(r1,primary,False);style_review_sheet(r2ws,r2,True)
    r2ws['G1']='Reviewer2 解剖标签（独立填写）';r2ws['H1']='Reviewer2 置信度（独立填写）';r2ws['I1']='Reviewer2 备注/判断依据'
    output.parent.mkdir(parents=True,exist_ok=True);wb.save(output)
    return output


def main():
    ap=argparse.ArgumentParser()
    ap.add_argument('--primary-csv',required=True)
    ap.add_argument('--double-review-manifest',required=True)
    ap.add_argument('--output',required=True)
    args=ap.parse_args()
    print(build(Path(args.primary_csv),Path(args.double_review_manifest),Path(args.output)))


if __name__=='__main__':
    main()
